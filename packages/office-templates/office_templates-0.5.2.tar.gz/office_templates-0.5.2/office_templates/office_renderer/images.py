from __future__ import annotations

from io import BytesIO
from typing import Callable, Optional
from urllib.request import urlopen

from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

from ..templating import process_text
from .constants import IMAGE_DIRECTIVES
from .exceptions import ImageError


def extract_image_directive(text: Optional[str]) -> tuple[Optional[str], Optional[str]]:
    """Return (url, mode) if *text* starts with an image directive."""
    if not text:
        return None, None
    stripped = text.strip()
    lowered = stripped.lower()
    for marker, mode in IMAGE_DIRECTIVES.items():
        if lowered.startswith(marker):
            url = stripped[len(marker) :].strip()
            return (url or None, mode)
    return None, None


def extract_image_url(text: Optional[str]) -> Optional[str]:
    """Backward compatible helper returning only the URL."""
    url, _ = extract_image_directive(text)
    return url


def should_replace_shape_with_image(shape) -> bool:
    """Return True if the shape text indicates an image placeholder."""
    if not hasattr(shape, "text_frame"):
        return False
    url, _ = extract_image_directive(shape.text_frame.text)
    return url is not None


def should_replace_cell_with_image(cell) -> bool:
    """Return True if the cell value indicates an image placeholder."""
    if not isinstance(cell.value, str):
        return False
    url, _ = extract_image_directive(cell.value)
    return url is not None


def replace_shape_with_image(
    shape,
    slide,
    context: dict,
    check_permissions: Optional[Callable[[object], bool]] = None,
    url: Optional[str] = None,
    mode: Optional[str] = None,
):
    """Replace *shape* with an image, keeping its position."""

    if url is None or mode is None:
        if not hasattr(shape, "text_frame"):
            return
        url, mode = extract_image_directive(shape.text_frame.text)
    if not url:
        return

    result = process_text(
        url,
        context=context,
        check_permissions=check_permissions,
        mode="normal",
    )
    assert isinstance(result, str), "Image URL must be a string"
    url = result

    try:
        with urlopen(url) as resp:
            data = resp.read()
    except Exception as e:  # pragma: no cover - network issues
        raise ImageError(f"Failed to download image from {url}: {e}")

    left = shape.left
    top = shape.top
    width = shape.width
    height = shape.height
    rotation = getattr(shape, "rotation", 0)

    if mode == "squeeze":
        pic = slide.shapes.add_picture(
            BytesIO(data), left, top, width=width, height=height
        )
    else:
        pic = slide.shapes.add_picture(BytesIO(data), left, top)
        with PILImage.open(BytesIO(data)) as pil_img:
            native_w, native_h = pil_img.size
        w_ratio = width / native_w
        h_ratio = height / native_h
        scale = min(w_ratio, h_ratio)
        pic.width = int(native_w * scale)
        pic.height = int(native_h * scale)
        pic.left = int(left + (width - pic.width) / 2)
        pic.top = int(top + (height - pic.height) / 2)

    pic.rotation = rotation

    # Remove the original shape
    sp_tree = shape._element.getparent()
    sp_tree.remove(shape._element)

    return pic


def replace_cell_with_image(
    cell,
    worksheet,
    context: dict,
    check_permissions: Optional[Callable[[object], bool]] = None,
    url: Optional[str] = None,
    mode: Optional[str] = None,
):
    """Replace the cell's value with an image anchored at the cell."""

    if url is None or mode is None:
        url, mode = extract_image_directive(
            cell.value if isinstance(cell.value, str) else None
        )
    if not url:
        return

    result = process_text(
        url,
        context=context,
        check_permissions=check_permissions,
        mode="normal",
    )
    assert isinstance(result, str), "Image URL must be a string"
    url = result

    try:
        with urlopen(url) as resp:
            data = resp.read()
    except Exception as e:  # pragma: no cover - network issues
        raise ImageError(f"Failed to download image from {url}: {e}")

    img = XLImage(BytesIO(data))
    img.anchor = cell.coordinate
    worksheet.add_image(img)
    cell.value = None
    return img
